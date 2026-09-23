"""
Reference BOQ parser — turn a human-priced SA bill workbook into typed ground truth.

Workbook conventions handled (observed on the Wedela BOQ, the common SA layout):

  Summary sheet        one row per building bill ('Existing Community Hall  Sum 1  R...')
                       + 'TOTAL Exl VAT' / 'VAT' / 'TOTAL INCL VAT'
  Building sheet       has a header row  ITEM NO | DESCRIPTION | UOM | QTY | RATE | SUB TOTAL
                       section rows      'A' | 'Distribution Boards / ...'
                       spec header rows  'A4.1' | '50mm2 x 4C PVC SWA ...'  (no qty)
                       child rows        None | 'Supply' | m | 50 | 938 | 46 900
                       section totals    'A- TOTAL Exl VAT' ... value in SUB TOTAL
                       roll-up block     'Summary' row, then per-section Sum lines,
                                         'Contigency @ 5%', 'TOTAL Exl VAT'
  P&Gs sheet           preliminaries (monthly site staff / establishment / plant)
  Rate sheets          'Installation Rate', 'Cable Termination material',
                       'Provisional Sum' — rate sources, not bill lines (ignored)

Rows with quantity 0/empty are not bill lines. Error cells ('#REF!') are recorded
on the building, never silently turned into numbers.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, Iterable, List, Optional

import openpyxl
from pydantic import BaseModel, Field

from evaluation.taxonomy import classify_item, line_role

_IGNORED_SHEETS = re.compile(r"installation rate|termination material|provisional sum", re.I)
_SECTION_TOTAL = re.compile(r"^\s*([A-Z])\s*-\s*TOTAL", re.I)
_ERROR_CELL = re.compile(r"^#(REF|VALUE|DIV/0|N/A|NAME|NUM)", re.I)


class RefLine(BaseModel):
    sheet: str
    building: str
    bill_section: str = ""          # A..D building bill, F/G/H kiosk, P prelims
    code: str = ""
    description: str
    parent: str = ""
    unit: str = ""
    qty: float
    rate: Optional[float] = None
    total: Optional[float] = None
    key_family: str
    key_spec: str = ""
    role: str = "combined"
    raw_row: int = 0

    @property
    def key(self) -> str:
        return f"{self.key_family}|{self.key_spec}" if self.key_spec else self.key_family

    @property
    def value(self) -> float:
        """Priced value of the line (0 when the bill left it unpriced)."""
        if self.total is not None:
            return float(self.total)
        return float(self.qty) * float(self.rate or 0.0)

    @property
    def label(self) -> str:
        """Human description, with the spec header for bare Supply/Install rows."""
        if self.parent and self.description.strip().lower() in ("supply", "install"):
            return f"{self.parent.strip()} — {self.description.strip()}"
        return self.description.strip()


class RefBuilding(BaseModel):
    name: str
    sheet: str
    lines: List[RefLine] = Field(default_factory=list)
    section_totals: Dict[str, float] = Field(default_factory=dict)
    contingency: float = 0.0
    total_excl_vat: float = 0.0
    in_summary: bool = False
    errors: List[str] = Field(default_factory=list)

    @property
    def value(self) -> float:
        return sum(l.value for l in self.lines)


class ReferenceBoq(BaseModel):
    project: str
    source_file: str = ""
    buildings: List[RefBuilding] = Field(default_factory=list)
    summary: Dict[str, float] = Field(default_factory=dict)   # summary-sheet name -> amount
    summary_total_excl_vat: float = 0.0
    vat: float = 0.0
    total_incl_vat: float = 0.0
    prelims: List[RefLine] = Field(default_factory=list)

    def building(self, name: str) -> Optional[RefBuilding]:
        want = _canon(name)
        return next((b for b in self.buildings if _canon(b.name) == want), None)

    def billed_buildings(self) -> List[RefBuilding]:
        return [b for b in self.buildings if b.in_summary]

    def all_lines(self) -> List[RefLine]:
        return [l for b in self.buildings for l in b.lines]


# ─── helpers ─────────────────────────────────────────────────────────

def _canon(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "")).strip().lower()


def _num(v) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _cells(row, n: int = 6) -> list:
    r = list(row[:n])
    return r + [None] * (n - len(r))


def _is_header(a, b) -> bool:
    return isinstance(a, str) and isinstance(b, str) and \
        a.strip().upper().startswith("ITEM") and b.strip().upper() == "DESCRIPTION"


def _has_bill_header(ws) -> bool:
    for row in ws.iter_rows(values_only=True, max_row=60):
        a, b = _cells(row, 2)
        if _is_header(a, b):
            return True
    return False


def _resolve_name(sheet: str, canonical: Iterable[str], summary_names: Iterable[str]) -> str:
    for pool in (canonical, summary_names):
        for n in pool:
            if _canon(n) == _canon(sheet):
                return n
    return sheet.strip()


# ─── sheet parsers ───────────────────────────────────────────────────

def _parse_summary(ws, ref: ReferenceBoq) -> None:
    for row in ws.iter_rows(values_only=True):
        a, b, _c, _d, _e, f = _cells(row)
        if isinstance(a, str):
            up = a.strip().upper()
            if up.startswith("TOTAL EXL") or up.startswith("TOTAL EXCL"):
                ref.summary_total_excl_vat = _num(f) or 0.0
            elif up.startswith("VAT"):
                ref.vat = _num(f) or 0.0
            elif up.startswith("TOTAL INCL"):
                ref.total_incl_vat = _num(f) or 0.0
            continue
        if isinstance(b, str) and _num(f) is not None:
            ref.summary[b.strip()] = _num(f)


def _parse_prelims(ws, ref: ReferenceBoq) -> None:
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        a, b, c, d, e, f = _cells(row)
        if _num(a) is None or not isinstance(b, str) or not _num(d):
            continue
        ref.prelims.append(RefLine(
            sheet=ws.title, building="P&Gs", bill_section="P", code=str(a),
            description=b.strip(), unit=str(c or ""), qty=_num(d), rate=_num(e),
            total=_num(f), key_family="prelims", role="combined", raw_row=i,
        ))


def _parse_building(ws, name: str) -> RefBuilding:
    bld = RefBuilding(name=name, sheet=ws.title)
    section, parent, rollup = "", "", False
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        a, b, c, d, e, f = _cells(row)

        for cell in (a, b, c, d, e, f):
            if isinstance(cell, str) and _ERROR_CELL.match(cell.strip()):
                label = a.strip() if isinstance(a, str) else f"row {i}"
                bld.errors.append(f"{cell.strip()} in '{label}' (row {i})")

        if isinstance(b, str) and b.strip().lower() == "summary":
            rollup = True
            continue
        if isinstance(a, str):
            m = _SECTION_TOTAL.match(a)
            if m:
                if _num(f) is not None:
                    bld.section_totals[m.group(1).upper()] = _num(f)
                continue
            up = a.strip().upper()
            if up.startswith("TOTAL EXL") or up.startswith("TOTAL EXCL"):
                bld.total_excl_vat = _num(f) or 0.0
                continue
            if up.startswith("VAT") or up.startswith("TOTAL INCL"):
                continue
        if rollup:
            if isinstance(b, str) and "contig" in b.lower():
                bld.contingency = _num(f) if _num(f) is not None else (_num(d) or 0.0)
            continue
        if _is_header(a, b):
            continue
        # section title row: single letter + title, no quantity
        if isinstance(a, str) and re.fullmatch(r"[A-Z]", a.strip()) and isinstance(b, str) and d is None:
            section, parent = a.strip(), ""
            continue
        if not isinstance(b, str):
            continue
        qty = _num(d)
        if qty is None:
            parent = b                      # spec header for following Supply/Install rows
            continue
        if qty == 0:
            continue
        key = classify_item(b, parent=parent, unit=str(c or ""))
        bld.lines.append(RefLine(
            sheet=ws.title, building=name, bill_section=section,
            code=str(a).strip() if a is not None else "", description=" ".join(b.split()),
            parent=" ".join(parent.split()), unit=str(c or "").strip(), qty=qty,
            rate=_num(e), total=_num(f), key_family=key.family, key_spec=key.spec,
            role=line_role(b), raw_row=i,
        ))
    return bld


# ─── entry point ─────────────────────────────────────────────────────

def parse_reference_xlsx(
    path: Path,
    *,
    project: str,
    buildings: Optional[List[str]] = None,
) -> ReferenceBoq:
    """
    Parse a priced BOQ workbook. `buildings` (e.g. from the dataset manifest)
    supplies canonical building names; sheet titles are matched to them
    case/space-insensitively.
    """
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    ref = ReferenceBoq(project=project, source_file=Path(path).name)

    for ws in wb.worksheets:
        if ws.title.strip().lower().startswith("summary"):
            _parse_summary(ws, ref)

    summary_canon = {_canon(n) for n in ref.summary}
    for ws in wb.worksheets:
        title = ws.title
        if title.strip().lower().startswith("summary") or _IGNORED_SHEETS.search(title):
            continue
        if re.search(r"p\s*&\s*gs?", title, re.I):
            _parse_prelims(ws, ref)
            continue
        if not _has_bill_header(ws):
            continue
        name = _resolve_name(title, buildings or [], ref.summary)
        bld = _parse_building(ws, name)
        bld.in_summary = _canon(name) in summary_canon or _canon(title) in summary_canon
        ref.buildings.append(bld)
    return ref


def load_reference(project: str) -> ReferenceBoq:
    """Load the committed parsed ground truth (data/projects/<p>/reference_boq.json)."""
    from evaluation.dataset import project_dir
    return ReferenceBoq.model_validate_json(
        (project_dir(project) / "reference_boq.json").read_text(encoding="utf-8")
    )
