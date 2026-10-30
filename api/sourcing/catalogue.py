"""
sourcing.catalogue — one dated price catalogue: supplier · item · price · price date.

Built from the price lists in doe/knowledge/price_lists/ (Schneider list prices), the AfriPlan
estimate prices in core.constants, and the supplier directory (Sparks buyer's guide). A user's
own prices live in a separate overrides CSV that a rebuild never touches: each override carries
its date and who made it, so the catalogue always says how fresh a price is.

Read-only like the rest of sourcing: no pipeline imports it. Deterministic, no LLM.
Build it with scripts/build_price_catalogue.py; read it with load_catalogue().
"""

from __future__ import annotations

import csv
import re
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

SCHNEIDER = "Schneider Electric"
AFRIPLAN = "AfriPlan estimate"


@dataclass
class CatalogueItem:
    item_id: str
    supplier: str
    brand: str
    part_no: str
    description: str
    family: str
    category: str = ""
    sub_category: str = ""
    discount_group: str = ""
    poles: Optional[int] = None
    amps: Optional[float] = None
    ka: Optional[float] = None
    unit: str = "No"
    list_price_zar_excl_vat: float = 0.0
    discount_pct: float = 0.0
    currency: str = "ZAR"
    price_date: str = ""
    price_source: str = "price list"
    source_document: str = ""
    afriplan_key: str = ""

    @property
    def net_price_zar_excl_vat(self) -> float:
        return round(self.list_price_zar_excl_vat * (1 - self.discount_pct / 100), 2)


@dataclass
class Supplier:
    name: str
    role: str            # manufacturer | distributor
    source: str
    source_date: str


@dataclass
class SourceDoc:
    document: str
    supplier: str
    price_date: str
    items: int
    skipped: int


COLUMNS = [f.name for f in fields(CatalogueItem)]
_INT_FIELDS = {"poles"}
_FLOAT_FIELDS = {"amps", "ka", "list_price_zar_excl_vat", "discount_pct"}


# ─── parsing helpers ─────────────────────────────────────────────────────────

def parse_zar(text: str) -> Optional[float]:
    """'ZAR 12 345,00' or '199.00' → float. Spaces inside the number are dropped (PDF extraction splits digits)."""
    s = re.sub(r"(?i)zar|r\s*(?=\d)", "", str(text or "")).replace(" ", "").replace(" ", "")
    if not re.fullmatch(r"\d+([.,]\d+)?", s):
        return None
    return float(s.replace(",", "."))


_POLES = re.compile(r"\b([1-4])\s?P\b", re.I)
_POLE_WORDS = {"SP": 1, "DP": 2, "TP": 3, "TPN": 4}
_POLE_WORD = re.compile(r"\b(SP|DP|TPN|TP)\b")
_AMPS = re.compile(r"(?<![\w.])(\d+(?:[.,]\d+)?)\s?A\b", re.I)
_KA = re.compile(r"(\d+(?:[.,]\d+)?)\s?kA\b", re.I)


def attributes(description: str) -> Dict[str, float]:
    """Poles, amps and kA written in a product description (only the ones present)."""
    out: Dict[str, float] = {}
    if m := _POLES.search(description):
        out["poles"] = int(m.group(1))
    elif m := _POLE_WORD.search(description):
        out["poles"] = _POLE_WORDS[m.group(1)]
    if m := _AMPS.search(description):
        out["amps"] = float(m.group(1).replace(",", "."))
    if m := _KA.search(description):
        out["ka"] = float(m.group(1).replace(",", "."))
    return out


# First match wins: specific families before generic ones.
_FAMILY_RULES: List[Tuple[str, re.Pattern]] = [(f, re.compile(p, re.I)) for f, p in [
    # industrial automation first, so its cables / switches are not taken for building items
    ("automation", r"automation|human[- ]machine|\bhmi\b|scada|telemetry|motion control|drives?\b|soft ?starter"),
    ("earth_leakage", r"rccb|rcbo|elcb|earth leakage|residual current|\d+\s?ma\b"),
    ("acb", r"air circuit|\bacb\b"),
    ("mccb", r"mccb|moulded case|ptccb|\bnsx|\bns\s?\d"),
    ("mcb", r"\bmcb\b|miniature circuit|\bi?c60"),
    ("surge", r"surge|\bspd\b"),
    ("isolator", r"isolator|switch[- ]?disconnector"),
    ("contactor", r"contactor"),
    ("socket", r"socket|outlet|receptacle"),
    ("switch", r"\bswitch|lever"),
    ("enclosure", r"enclosure|distribution board|\bdb\b|\bmodules?\b|kaedra"),
    ("busbar", r"busbar|comb bar"),
    ("cable", r"cable|\bwire\b|\bswa\b|surfix|\bgp\b"),
    ("luminaire", r"luminaire|downlight|light fitting|\bled\b|lamp|vapou?r|prismatic|bulkhead|flood|fluorescent"),
    ("containment", r"conduit|trunking|skirting|tray|sleeve|\bbox\b"),
    ("control", r"contact|push ?button|pilot|signalling|selector"),
]]


def classify_family(description: str, category: str = "") -> str:
    text = f"{description} {category}"
    for family, pat in _FAMILY_RULES:
        if pat.search(text):
            return family
    return "other"


def _clean(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _item(**kw) -> CatalogueItem:
    attrs = attributes(kw["description"])
    kw.setdefault("family", classify_family(kw["description"], kw.get("category", "")))
    return CatalogueItem(poles=attrs.get("poles"), amps=attrs.get("amps"), ka=attrs.get("ka"), **kw)


# ─── sources ─────────────────────────────────────────────────────────────────

def parse_schneider_rows(rows: Iterable[List[str]], *, source: str, price_date: str,
                         stats: Optional[Dict[str, int]] = None) -> Tuple[List[CatalogueItem], int]:
    """Rows of the Schneider list table → items. Returns (items, rows skipped for having no price).

    A part number seen again with the same description and price is an exact repeat (dropped and
    counted); with a different price or range it is kept as `<id>~2`, so nothing disappears silently.
    """
    items: List[CatalogueItem] = []
    seen: Dict[str, List[CatalogueItem]] = {}
    skipped = 0
    counts = {"exact_duplicates": 0, "same_part_other_price": 0}
    for r in rows:
        r = [_clean(c) for c in r] + [""] * (12 - len(r))
        if r[0] == "Collection" or r[4] == "Part" or not r[4] or " " in r[4]:
            continue
        price = parse_zar(r[6])
        if price is None:
            skipped += 1 if r[5] else 0
            continue
        twins = seen.setdefault(r[4], [])
        if any(t.description == r[5] and t.list_price_zar_excl_vat == price for t in twins):
            counts["exact_duplicates"] += 1
            continue
        if twins:
            counts["same_part_other_price"] += 1
        item_id = f"schneider:{r[4]}" + (f"~{len(twins) + 1}" if twins else "")
        item = _item(
            item_id=item_id, supplier=SCHNEIDER, brand=SCHNEIDER, part_no=r[4],
            description=r[5], category=r[2], sub_category=r[1], discount_group=r[3],
            list_price_zar_excl_vat=price, price_date=price_date, source_document=source,
        )
        twins.append(item)
        items.append(item)
    if stats is not None:
        stats.update(counts)
    return items, skipped


Span = Tuple[float, List[Tuple[float, str]]]   # (x0, [(char centre x, char)])
_PART, _DESC = 4, 5


def row_cells(spans: List[Span], edges: List[float]) -> List[str]:
    """One table row → cells. A span belongs to the column it starts in. It is split at a later column
    edge only when that column has no span of its own on the row: neighbouring spreadsheet cells are
    sometimes printed as one glued span ("…General PEntry & Medium Meters"), while a description that
    runs on into the price column stays whole (the price has its own span)."""
    cells = [""] * len(edges)

    def col_of(x: float) -> int:
        return max([i for i, e in enumerate(edges) if x >= e - 2] or [0])

    starts = {col_of(x0) for x0, _ in spans}
    for x0, chars in sorted(spans, key=lambda s: s[0]):
        col = col_of(x0)
        split_at = [c for c in range(col + 1, len(edges)) if c not in starts]
        parts: Dict[int, str] = {}
        for xc, ch in chars:
            k = max([c for c in split_at if xc >= edges[c]] or [col])
            if k != col and any(c in starts for c in range(col + 1, k + 1)):
                k = col  # never jump over a column that has its own span
            parts[k] = parts.get(k, "") + ch
        for k, text in parts.items():
            cells[k] += ("" if not cells[k] or cells[k].endswith(" ") else " ") + text
    return [_clean(c) for c in cells]


_HEADER_WORDS = ["Collection", "Sub-collection", "Local", "Discount", "Part", "Description",
                 "List", "Activity", "Length", "Height", "Width", "Weight"]


def read_schneider_pdf(path: Path) -> Tuple[List[List[str]], str]:
    """All table rows of the Schneider price-list PDF, and its date (from the PDF metadata).

    Reads text spans and places them by the header's column edges (~0.1 s a page; table detection
    took ~3 s a page and pushed overflowing description letters into the price cell)."""
    import fitz  # PyMuPDF — already an api dependency

    doc = fitz.open(path)
    edges: List[float] = []
    for page in doc:
        words = {w[4]: w[0] for w in reversed(page.get_text("words"))}
        if all(h in words for h in _HEADER_WORDS):
            edges = [words[h] for h in _HEADER_WORDS]
            break
    if not edges:
        raise ValueError(f"{path}: no price-list header row ({', '.join(_HEADER_WORDS)}) found")

    rows: List[List[str]] = []
    for page in doc:
        by_line: Dict[int, List[Span]] = {}
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                for sp in line["spans"]:
                    chars = [((c["bbox"][0] + c["bbox"][2]) / 2, c["c"]) for c in sp["chars"]]
                    if chars:
                        by_line.setdefault(round(sp["bbox"][1] / 2), []).append((sp["bbox"][0], chars))
        rows.extend(row_cells(spans, edges) for _, spans in sorted(by_line.items()))
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", doc.metadata.get("creationDate") or "")
    return rows, (f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else "")


_CONSTANT_TABLES = {  # core.constants map → default family when the key alone does not say
    "LIGHT_PRICES": "luminaire", "SOCKET_PRICES": "socket", "SWITCH_PRICES": "switch",
    "CABLE_PRICES": "cable", "DB_PRICES": "enclosure", "CIVIL_PRICES": "civil",
    "CONTAINMENT_PRICES": "containment",
}  # LABOUR_RATES left out: labour is not a product


def items_from_constants(*, price_date: str) -> List[CatalogueItem]:
    """The AfriPlan estimate prices (core.constants) as catalogue rows, keyed by their constant name."""
    from core import constants

    items: List[CatalogueItem] = []
    for table, default_family in _CONSTANT_TABLES.items():
        for key, price in getattr(constants, table).items():
            desc = key.replace("_", " ")
            family = classify_family(desc)
            items.append(_item(
                item_id=f"afriplan:{key}", supplier=AFRIPLAN, brand="", part_no=key, description=desc,
                family=family if family != "other" else default_family, category=table,
                list_price_zar_excl_vat=float(price), price_date=price_date,
                price_source="AfriPlan estimate", source_document="api/core/constants.py", afriplan_key=key,
            ))
    return items


_SPARKS_HEADINGS = {
    "BUYER'S GUIDE", "CONTACTS", "PROTECTION", "FUSES", "CONTACTORS", "METERS", "TIMERS", "BREAKERS",
    "SPARKS", "ELECTRICAL NEWS", "SEPTEMBER 2019", "LABELLING", "HAZARDOUS AREAS", "SWITCHES/ISOLATORS",
    "SURGE PROTECTION DEVICES", "DISTRIBUTION BOARDS", "DISTRIBUTION BOARD ENCLOSURES",
    "DISTRIBUTION BOARDS, SWITCHES, SOCKETS AND PROTECTION", "CIRCUIT BREAKERS / EARTH LEAKAGE",
    "(TTA/MCC)", "40-1600 A",
}


def parse_sparks_directory(text: str, *, source: str, source_date: str) -> List[Supplier]:
    """Company names under MANUFACTURERS / DISTRIBUTORS. Stops at CONTACT: people's names are not kept."""
    role = ""
    seen: Dict[Tuple[str, str], Supplier] = {}
    for raw in text.splitlines():
        line = _clean(raw)
        if line == "MANUFACTURERS":
            role = "manufacturer"
        elif line == "DISTRIBUTORS":
            role = "distributor"
        elif line == "CONTACT":
            break
        elif role and line and line not in _SPARKS_HEADINGS and not re.fullmatch(r"[\d +()\-]+", line):
            seen.setdefault((line, role), Supplier(line, role, source, source_date))
    return list(seen.values())


# ─── user overrides ──────────────────────────────────────────────────────────

def apply_overrides(items: List[CatalogueItem], overrides: Iterable[Dict[str, str]]) -> Tuple[List[CatalogueItem], List[str]]:
    """A user's prices on top of the built catalogue. A row for an unknown item with a description adds it."""
    by_id = {i.item_id: i for i in items}
    problems: List[str] = []
    for o in overrides:
        o = {k: _clean(v) for k, v in o.items() if k}
        item_id = o.get("item_id", "")
        if not item_id:
            continue
        who = o.get("updated_by") or "user"
        source = f"override: {who}" + (f" — {o['note']}" if o.get("note") else "")
        price = parse_zar(o.get("list_price_zar_excl_vat", "")) if o.get("list_price_zar_excl_vat") else None
        item = by_id.get(item_id)
        if item is None:
            if not o.get("description"):
                problems.append(f"{item_id}: not in the catalogue and no description — ignored")
                continue
            item = _item(item_id=item_id, supplier=o.get("supplier") or item_id.split(":")[0],
                         brand=o.get("brand", ""), part_no=o.get("part_no") or item_id.split(":")[-1],
                         description=o["description"], unit=o.get("unit") or "No",
                         category=o.get("category", ""), source_document="price_overrides.csv")
            by_id[item_id] = item
        if price is not None:
            item.list_price_zar_excl_vat = price
        if o.get("discount_pct"):
            item.discount_pct = float(o["discount_pct"].replace(",", "."))
        if o.get("price_date"):
            item.price_date = o["price_date"]
        item.price_source = source
    return list(by_id.values()), problems


def read_overrides(path: Path) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


OVERRIDE_COLUMNS = ["item_id", "list_price_zar_excl_vat", "discount_pct", "price_date", "updated_by", "note",
                    "supplier", "description", "unit", "category"]


# ─── writers / reader ────────────────────────────────────────────────────────

def write_csv(items: List[CatalogueItem], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for i in items:
            w.writerow({k: ("" if v is None else v) for k, v in asdict(i).items()})


def load_catalogue(path: Path) -> List[CatalogueItem]:
    """The catalogue CSV as items (the web app's entry point)."""
    out: List[CatalogueItem] = []
    with Path(path).open(newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            kw: Dict[str, object] = {}
            for k in COLUMNS:
                v = row.get(k, "")
                if k in _INT_FIELDS:
                    kw[k] = int(v) if v else None
                elif k in _FLOAT_FIELDS:
                    kw[k] = float(v) if v else (0.0 if k in ("list_price_zar_excl_vat", "discount_pct") else None)
                else:
                    kw[k] = v
            out.append(CatalogueItem(**kw))
    return out


_README = [
    "AfriPlan price catalogue — one row per product: supplier, description, family, list price excl. VAT, price date.",
    "Edit prices in price_overrides.csv (item_id, new price, price_date, updated_by, note), then rebuild:",
    "    <venv python> scripts/build_price_catalogue.py",
    "A rebuild re-reads the price lists and applies the overrides on top; overrides are never lost.",
    "discount_pct is your trade discount; net_price = list × (1 − discount/100). Prices exclude VAT.",
    "price_source says where each price comes from: 'price list', 'AfriPlan estimate' or 'override: <who>'.",
    "Old prices: check price_date before tendering — list prices change every year.",
]


def write_xlsx(items: List[CatalogueItem], suppliers: List[Supplier], sources: List[SourceDoc], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    for line in _README:
        readme.append([line])
    readme.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Catalogue")
    head = COLUMNS + ["net_price_zar_excl_vat"]
    ws.append(head)
    lp = get_column_letter(COLUMNS.index("list_price_zar_excl_vat") + 1)
    dc = get_column_letter(COLUMNS.index("discount_pct") + 1)
    for r, i in enumerate(items, start=2):
        ws.append([getattr(i, k) for k in COLUMNS] + [f"=ROUND({lp}{r}*(1-{dc}{r}/100),2)"])
    widths = {"item_id": 22, "supplier": 20, "description": 48, "family": 14, "category": 30, "price_source": 22}
    for c, name in enumerate(head, start=1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for title, rows, cols in (("Suppliers", suppliers, ["name", "role", "source", "source_date"]),
                              ("Sources", sources, ["document", "supplier", "price_date", "items", "skipped"])):
        sh = wb.create_sheet(title)
        sh.append(cols)
        for x in rows:
            sh.append([getattr(x, k) for k in cols])
        for c in range(1, len(cols) + 1):
            sh.cell(row=1, column=c).font = Font(bold=True)
            sh.column_dimensions[get_column_letter(c)].width = 34
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
