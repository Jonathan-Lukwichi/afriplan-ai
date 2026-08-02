"""
Tests for the ported tender-document exporters (exports/excel_boq.py,
exports/pdf_boq.py) - both copied verbatim from the original app, which had
no dedicated tests for them (its own manual click-through was the only
verification). These are new, not ported.
"""

from __future__ import annotations

from openpyxl import load_workbook

from agent.shared import BillOfQuantities, BQLineItem, BQSection, ContractorProfile, ItemConfidence, ProjectMetadata
from exports import export_boq_to_excel, export_boq_to_pdf


def _make_boq() -> BillOfQuantities:
    items = [
        BQLineItem(item_no=1, section=BQSection.LIGHTING, description="LED Downlight",
                   qty=24, unit_price_zar=220.0, total_zar=5280.0, source=ItemConfidence.EXTRACTED),
        BQLineItem(item_no=2, section=BQSection.POWER_OUTLETS, description="Double Socket Outlet",
                   qty=20, unit_price_zar=160.0, total_zar=3200.0, source=ItemConfidence.EXTRACTED),
    ]
    subtotal = sum(i.total_zar for i in items)
    return BillOfQuantities(
        project_name="Test Project", pipeline="dxf", run_id="abc123456789",
        line_items=items, subtotal_zar=subtotal,
        total_excl_vat_zar=round(subtotal * 1.25, 2),
        total_incl_vat_zar=round(subtotal * 1.25 * 1.15, 2),
        vat_zar=round(subtotal * 1.25 * 0.15, 2),
        contingency_zar=round(subtotal * 0.05, 2),
        markup_zar=round(subtotal * 0.20, 2),
        items_extracted=len(items),
    )


def test_excel_export_produces_valid_workbook():
    boq = _make_boq()
    xlsx_bytes = export_boq_to_excel(
        boq, project=ProjectMetadata(project_name="Test Project"),
        contractor=ContractorProfile(), quote_ref="AFP-TEST-0001", validity_days=30,
    )
    assert xlsx_bytes.startswith(b"PK")  # xlsx is a zip container
    assert len(xlsx_bytes) > 1000

    import io
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert len(wb.sheetnames) >= 1


def test_pdf_export_produces_valid_pdf():
    boq = _make_boq()
    pdf_bytes = export_boq_to_pdf(
        boq, project=ProjectMetadata(project_name="Test Project"),
        contractor=ContractorProfile(), quote_ref="AFP-TEST-0001", validity_days=30,
    )
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000


def test_exports_use_defaults_when_project_and_contractor_omitted():
    boq = _make_boq()
    xlsx_bytes = export_boq_to_excel(boq)
    pdf_bytes = export_boq_to_pdf(boq)
    assert xlsx_bytes.startswith(b"PK")
    assert pdf_bytes.startswith(b"%PDF")
