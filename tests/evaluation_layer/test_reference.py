"""Reference BOQ parser — the human-priced bill as typed ground truth."""

from pathlib import Path

import openpyxl
import pytest

from evaluation.dataset import project_dir, raw_available
from evaluation.reference import parse_reference_xlsx

HDR = ("ITEM NO", "DESCRIPTION", "UOM", "QTY", "RATE", "SUB TOTAL")


def _synthetic(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Summary  Test BOQ"
    s.append(("ITEM No", "DESCRIPTION", "UoM", "QTY", "Rate", "Sub Total"))
    s.append((1, "Hall", "Sum", 1, 1050.0, 1050.0))
    s.append(("TOTAL Exl VAT", None, None, None, None, 1050.0))
    s.append(("VAT ( @ 15% )", None, None, None, None, 157.5))
    s.append(("TOTAL INCL VAT", None, None, None, None, 1207.5))

    h = wb.create_sheet("Hall ")
    h.append(("A", "Distribution Boards / Excavations/ Main cables"))
    h.append(HDR)
    h.append(("EXCAVATIONS",))
    h.append(("A1.1", "Trenching and re-instatement", "m", 10, 50.0, 500.0))
    h.append(("A1.3", "110 mm diameter flexible PVC cable sleeves.", "m", 0, 149.8, 0))
    h.append(("A4.1", "16mm2 x 4C PVC SWA PVC 600-1000V Cable"))
    h.append(("A4.2", "Supply", "m", 10, 30.0, 300.0))
    h.append((None, "Install", "m", 10, 20.0, 200.0))
    h.append(("A- TOTAL Exl VAT", None, None, None, None, 1000.0))
    h.append(("B", "General lighting"))
    h.append(HDR)
    h.append(("D1.1", "30watt led flood light", "No", 2, None, None))
    h.append(("B-TOTAL Exl VAT", None, None, None, None, "#REF!"))
    h.append((None, "Summary"))
    h.append(HDR)
    h.append(("A", "Distribution Boards", "Sum", 1, 1000.0, 1000.0))
    h.append(("E", "Contigency @ 5%", "Sum", 50.0, None, 50.0))
    h.append(("TOTAL Exl VAT", None, None, None, None, 1050.0))

    x = wb.create_sheet("Spare Block")
    x.append(("A", "Distribution Boards"))
    x.append(HDR)
    x.append(("A1.1", "Trenching", "m", 5, 50.0, 250.0))
    x.append(("TOTAL Exl VAT", None, None, None, None, 250.0))

    p = tmp_path / "boq.xlsx"
    wb.save(p)
    return p


def test_synthetic_workbook(tmp_path):
    ref = parse_reference_xlsx(_synthetic(tmp_path), project="t", buildings=["Hall"])
    assert ref.summary_total_excl_vat == 1050.0
    hall = ref.building("Hall")
    assert hall is not None and hall.in_summary
    descs = [(l.bill_section, l.key_family, l.key_spec, l.role, l.qty) for l in hall.lines]
    assert ("A", "trench", "", "combined", 10) in descs
    assert ("A", "swa_cable", "16mm2|4c", "supply", 10) in descs
    assert ("A", "swa_cable", "16mm2|4c", "install", 10) in descs
    assert all(l.key_family != "sleeve" for l in hall.lines)          # qty 0 dropped
    flood = [l for l in hall.lines if l.key_family == "light_flood"][0]
    assert flood.bill_section == "B" and flood.rate is None and flood.value == 0.0
    assert hall.section_totals["A"] == 1000.0
    assert hall.contingency == 50.0 and hall.total_excl_vat == 1050.0
    assert any("#REF!" in e for e in hall.errors)
    # summary rows (A 'Distribution Boards' Sum 1) are NOT bill lines
    assert all(l.unit != "Sum" or l.key_family != "other" for l in hall.lines)

    spare = ref.building("Spare Block")
    assert spare is not None and not spare.in_summary


@pytest.mark.skipif(not raw_available("wedela"), reason="Wedela raw files not present")
def test_real_wedela_workbook():
    path = project_dir("wedela") / "raw" / "Wedela BOQ Rev01 141125.xlsx"
    ref = parse_reference_xlsx(path, project="wedela")
    assert ref.summary_total_excl_vat == pytest.approx(sum(ref.summary.values()), abs=1)
    in_summary = sorted(b.name for b in ref.buildings if b.in_summary)
    assert len(in_summary) == 7
    heat = ref.building("Pool-Heat Pumps")
    assert heat is not None and not heat.in_summary and heat.errors
    ech = ref.building("Existing Community Hall")
    swa = [l for l in ech.lines if l.key_family == "swa_cable" and l.role == "supply"]
    assert swa and swa[0].key_spec == "50mm2|4c" and swa[0].qty == 50
    assert ech.total_excl_vat == pytest.approx(ref.summary["Existing Community Hall"], abs=1)
    assert len(ref.prelims) >= 16 and sum(l.value for l in ref.prelims) == pytest.approx(ref.summary["P&Gs"])
    # every priced line of every billed building is classified (no silent 'other' drift)
    others = [l.description for b in ref.buildings if b.in_summary for l in b.lines
              if l.key_family == "other" and l.value > 0]
    assert len(others) <= 3, others
